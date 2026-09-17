#!/usr/bin/env python3
"""Batch Dynamics_predicter over Zenodo Media anchors for all instruments/effects.

Matches the GUI defaults in start.bat / run_gui.py: PCHIP on, n_boot=200.
Writes one 10-dynamic workbook per technique into that instrument folder
(Dynamics10 subfolder). Does not overwrite Zenodo/STE source books.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from dynamics_predicter.transfer import (  # noqa: E402
    ColumnMap,
    __version__,
    export_excel,
    load_panel_xlsx,
    outer_sensitivity_table,
    run_pipeline,
    validate_panel_df,
)

DESKTOP = Path(r"C:\Users\lmr20\Desktop\Código extrapolação")

JOBS = [
    # instrument_id, folder, prefix, source xlsx, technique_slug
    ("violin", "VIOLIN_3", "Violin", "VIOLIN_Zenodo_collections_Arco_normal.xlsx", "Arco_normal"),
    ("violin", "VIOLIN_3", "Violin", "Violin_Zenodo_collections_con_sordino.xlsx", "con_sordino"),
    ("violin", "VIOLIN_3", "Violin", "Violin_Zenodo_collections_sul_ponticello.xlsx", "sul_ponticello"),
    ("violin", "VIOLIN_3", "Violin", "Violin_Zenodo_collections_sul_tasto.xlsx", "sul_tasto"),
    ("violin", "VIOLIN_3", "Violin", "Violin_Zenodo_collections_harmonics.xlsx", "harmonics"),
    ("viola", "VIOLA", "Viola", "VIOLA_Zenodo_collections_Arco_normal.xlsx", "Arco_normal"),
    ("viola", "VIOLA", "Viola", "Viola_Zenodo_collections_con_sordino.xlsx", "con_sordino"),
    ("viola", "VIOLA", "Viola", "Viola_Zenodo_collections_sul_ponticello.xlsx", "sul_ponticello"),
    ("viola", "VIOLA", "Viola", "Viola_Zenodo_collections_harmonics.xlsx", "harmonics"),
    ("cello", "CELLO", "Cello", "CELLO_Zenodo_collections_media.xlsx", "Arco_normal"),
    ("cello", "CELLO", "Cello", "Cello_Zenodo_collections_con_sordino.xlsx", "con_sordino"),
    ("cello", "CELLO", "Cello", "Cello_Zenodo_collections_sul_ponticello.xlsx", "sul_ponticello"),
    ("cello", "CELLO", "Cello", "Cello_Zenodo_collections_harmonics.xlsx", "harmonics"),
    ("double_bass", "DOUBLE_BASS", "DoubleBass", "DOUBLEBASS_Zenodo_collections_media.xlsx", "Arco_normal"),
    ("double_bass", "DOUBLE_BASS", "DoubleBass", "DoubleBass_Zenodo_collections_con_sordino.xlsx", "con_sordino"),
    ("double_bass", "DOUBLE_BASS", "DoubleBass", "DoubleBass_Zenodo_collections_sul_ponticello.xlsx", "sul_ponticello"),
    ("double_bass", "DOUBLE_BASS", "DoubleBass", "DoubleBass_Zenodo_collections_harmonics.xlsx", "harmonics"),
]


def _norm(h) -> str:
    if h is None:
        return ""
    return re.sub(r"\s+", "_", str(h).strip().lower())


def pick_media_sheet(sheet_names: list[str]) -> str:
    skip = {"media_pp", "media_mf", "media_ff"}
    cands = [
        s
        for s in sheet_names
        if _norm(s) not in skip and ("media" in _norm(s))
    ]
    if not cands:
        raise ValueError(f"No Media sheet in {sheet_names}")
    # Prefer *Media / *_Media over Media_pp-style leftovers
    preferred = [s for s in cands if _norm(s).endswith("_media") or _norm(s).endswith("media")]
    return preferred[0] if preferred else cands[0]


def find_header_index(norms: list[str], *candidates: str) -> int | None:
    for cand in candidates:
        hits = [i for i, h in enumerate(norms) if h == cand]
        if hits:
            return hits[-1]  # compact Media block is to the right of the long average headers
    return None


def _header_index(header: tuple, *names: str) -> int | None:
    norms = [_norm(c) for c in header]
    want = {_norm(n) for n in names}
    for i, h in enumerate(norms):
        if h in want:
            return i
    return None


def reconstruct_media_from_dyn_sheets(path: Path) -> pd.DataFrame:
    """Rebuild Media pp/mf/ff when the Media sheet formulas were not cached.

    Media_d = nanmean(IOWA CDM, ORCH CDM) for each note — same as the
    Zenodo Media formula (IOWA + ORCH) / 2 when both exist.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    by_note: dict[str, dict[str, list[float]]] = {}
    used: list[str] = []
    for sheet in wb.sheetnames:
        sl = sheet.lower().replace("__", "_")
        dyn = None
        coll = None
        for d in ("pp", "mf", "ff"):
            if sl.endswith("_" + d):
                dyn = d
                break
        if dyn is None:
            continue
        if "iowa" in sl:
            coll = "iowa"
        elif "orch" in sl:
            coll = "orch"
        else:
            continue
        if "empirical" in sl or "media" in sl:
            continue
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = rows[0]
        note_i = _header_index(header, "Notes", "Note", "Source note", "notas")
        val_i = _header_index(
            header,
            "CDM - Media",
            "Combined density metric",
            "CDM",
        )
        if note_i is None or val_i is None:
            continue
        n_got = 0
        for row in rows[1:]:
            if row is None or len(row) <= max(note_i, val_i):
                continue
            note = row[note_i]
            if note is None or (isinstance(note, (int, float)) and not isinstance(note, bool)):
                continue
            note_s = str(note).strip()
            if not note_s:
                continue
            try:
                val = float(row[val_i])
            except (TypeError, ValueError):
                continue
            if val <= 0:
                continue
            by_note.setdefault(note_s, {}).setdefault(dyn, []).append(val)
            n_got += 1
        if n_got:
            used.append(f"{sheet}:{n_got}")
    wb.close()
    records = []
    for note, dyns in by_note.items():
        rec = {
            "note": note,
            "note_target": note,
            "tgt_pp": None,
            "tgt_mf": None,
            "tgt_ff": None,
        }
        for d in ("pp", "mf", "ff"):
            vals = dyns.get(d) or []
            if vals:
                rec[f"tgt_{d}"] = float(sum(vals) / len(vals))
        rec["complete_anchors"] = all(rec[f"tgt_{d}"] is not None for d in ("pp", "mf", "ff"))
        records.append(rec)
    df = pd.DataFrame(records)
    if df.empty:
        raise ValueError(f"{path.name}: dyn-sheet reconstruct produced 0 rows (used={used})")
    hard = [e for e in validate_panel_df(df) if not e.startswith("duplicate")]
    if hard:
        raise ValueError(f"{path.name}: reconstructed panel invalid: " + "; ".join(hard))
    df.attrs["reconstruct_sheets"] = used
    return df


def column_map_for_media(path: Path) -> tuple[str, ColumnMap]:
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet = pick_media_sheet(list(wb.sheetnames))
    ws = wb[sheet]
    header = next(ws.iter_rows(max_row=1, values_only=True))
    wb.close()
    norms = [_norm(c) for c in header]

    note = find_header_index(norms, "note", "notas", "notes")
    if note is None:
        raise ValueError(f"{path.name}: no Note column on {sheet}")

    pp = find_header_index(
        norms,
        "media_pp",
        "cross-collection_mean_pp",
        "cross_collection_mean_pp",
    )
    mf = find_header_index(
        norms,
        "media_mf",
        "cross-collection_mean_mf",
        "cross_collection_mean_mf",
    )
    ff = find_header_index(
        norms,
        "media_ff",
        "cross-collection_mean_ff",
        "cross_collection_mean_ff",
    )
    if pp is None or mf is None or ff is None:
        raise ValueError(
            f"{path.name} sheet {sheet}: could not find Media pp/mf/ff "
            f"(pp={pp} mf={mf} ff={ff}); headers={header[:20]}"
        )
    cmap = ColumnMap(
        note=note,
        pp=pp,
        mf=mf,
        ff=ff,
        note_target=None,
        data_start_row=1,
        sheet=sheet,
    )
    return sheet, cmap


def run_one(job: tuple[str, str, str, str, str], *, n_boot: int, dry_run: bool) -> dict:
    instrument_id, folder, prefix, fname, technique = job
    src = DESKTOP / folder / fname
    out_dir = DESKTOP / folder / "Dynamics10"
    out = out_dir / f"{prefix}_Dynamics10_{technique}.xlsx"
    rec: dict = {
        "instrument": instrument_id,
        "technique": technique,
        "source": str(src),
        "out": str(out),
    }
    if not src.is_file():
        rec["ok"] = False
        rec["error"] = f"missing source {src}"
        return rec

    sheet, cmap = column_map_for_media(src)
    rec["cols"] = cmap.to_dict()
    rec["sheet"] = sheet
    try:
        df = load_panel_xlsx(src, column_map=cmap)
        rec["source_mode"] = "media_sheet"
    except ValueError:
        df = reconstruct_media_from_dyn_sheets(src)
        rec["source_mode"] = "reconstructed_dyn_sheets"
        rec["reconstruct_sheets"] = df.attrs.get("reconstruct_sheets")
    n_complete = int(df["complete_anchors"].sum()) if "complete_anchors" in df.columns else len(df)
    rec["n_rows"] = int(len(df))
    rec["n_complete"] = n_complete
    if n_complete == 0:
        df = reconstruct_media_from_dyn_sheets(src)
        rec["source_mode"] = "reconstructed_dyn_sheets"
        rec["reconstruct_sheets"] = df.attrs.get("reconstruct_sheets")
        n_complete = int(df["complete_anchors"].sum())
        rec["n_rows"] = int(len(df))
        rec["n_complete"] = n_complete
    if n_complete == 0:
        rec["ok"] = False
        rec["error"] = "0 complete pp/mf/ff rows"
        return rec

    work = df.loc[df["complete_anchors"]].copy() if "complete_anchors" in df.columns else df
    if dry_run:
        rec["ok"] = True
        rec["dry_run"] = True
        return rec

    out_dir.mkdir(parents=True, exist_ok=True)
    results, holdouts, _bootstrap, _loo, meta = run_pipeline(
        work,
        source=str(src),
        use_pchip=True,
        n_boot=n_boot,
    )
    sens = outer_sensitivity_table(results, work)
    export_excel(
        results,
        holdouts,
        out,
        meta,
        sensitivity_df=sens,
        calibration=meta.get("span_model_calibration_loo"),
    )
    rec["ok"] = True
    rec["n_ok"] = meta.get("n_ok")
    rec["mean_quality"] = meta.get("mean_quality")
    rec["n_nonmonotonic_anchors"] = meta.get("n_nonmonotonic_anchors")
    rec["version"] = __version__
    return rec


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--instrument", action="append", help="violin / viola / cello / double_bass")
    p.add_argument("--n-boot", type=int, default=200)
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args(argv)

    jobs = JOBS
    if args.instrument:
        want = {x.lower().replace("-", "_") for x in args.instrument}
        jobs = [j for j in JOBS if j[0] in want]
        if not jobs:
            print("No jobs matched", want)
            return 2

    summaries = []
    n_ok = 0
    for i, job in enumerate(jobs, 1):
        label = f"{job[0]}/{job[4]}"
        print(f"[{i}/{len(jobs)}] {label} …", flush=True)
        try:
            rec = run_one(job, n_boot=args.n_boot, dry_run=args.dry_run)
        except Exception as exc:  # noqa: BLE001
            rec = {
                "instrument": job[0],
                "technique": job[4],
                "ok": False,
                "error": str(exc),
                "traceback": traceback.format_exc(),
            }
        summaries.append(rec)
        if rec.get("ok"):
            n_ok += 1
            print(
                f"  OK n={rec.get('n_complete')} sheet={rec.get('sheet')} -> {rec.get('out')}",
                flush=True,
            )
        else:
            print(f"  FAIL {rec.get('error')}", flush=True)

    log_path = DESKTOP / "Dynamics10_batch_log.json"
    log_path.write_text(json.dumps(summaries, indent=2), encoding="utf-8")
    print(f"Done {n_ok}/{len(jobs)}. Log: {log_path}")
    return 0 if n_ok == len(jobs) else 1


if __name__ == "__main__":
    raise SystemExit(main())
